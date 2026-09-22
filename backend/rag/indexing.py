"""Document ingestion: text extraction, chunking, and the per-user FAISS index."""
import csv
import os
import pickle
import re

import faiss

from rag.embeddings import EMBED_MODEL_NAME, embed_query, get_embed_model
from rag.ocr import IMAGE_EXTENSIONS, OCRError, extract_image_text, extract_pdf_text
from rag.tracing import (
    TRACE_APP_TAG,
    _attributes,
    _emit_event,
    _flush,
    _get_langfuse_client,
    _observe,
    _tag_trace,
    _update,
)

CHUNK_SIZE = 500        # characters per chunk
CHUNK_OVERLAP = 50
VECTORSTORE_DIR = os.path.join(os.path.dirname(__file__), "..", "vectorstore")


def _user_has_index(user_id):
    """True agar is user ne koi document upload kar rakha hai."""
    index_path, _ = _paths_for_user(user_id)
    return os.path.exists(index_path)


def _extract_docx_text(file_path):
    """Word file: paragraphs + tables (table cells ek line me '|' se juda)."""
    from docx import Document

    doc = Document(file_path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_csv_text(file_path):
    """CSV: har row 'column: value, column: value' line ban jati hai -- retrieval ke liye behtar."""
    with open(file_path, "r", encoding="utf-8-sig", errors="ignore", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel
        rows = [r for r in csv.reader(f, dialect) if any(c.strip() for c in r)]
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    if not body:
        return ", ".join(header)
    return "\n".join(
        ", ".join(f"{h.strip() or f'col{i + 1}'}: {v.strip()}" for i, (h, v) in enumerate(zip(header, row)))
        for row in body
    )


def extract_text_ex(file_path):
    """
    File -> (text, info). `info` me OCR ki tafseel hoti hai (kitne pages OCR
    hue, kitne fail hue) jo Langfuse ke extract-text span me dikhti hai.

    Scanned PDFs aur images (Chinese/complex bhi) rag/ocr.py se OCR hote hain;
    baaki formats direct parse hote hain. OCR ki koi user-facing ghalti
    OCRError ban kar upar jati hai (app.py /upload usay saaf message me badalta hai).
    """
    ext = os.path.splitext(file_path)[1].lower().lstrip(".")
    if ext == "pdf":
        return extract_pdf_text(file_path)
    if ext in IMAGE_EXTENSIONS:
        return extract_image_text(file_path)
    if ext == "docx":
        return _extract_docx_text(file_path), {"kind": "docx"}
    if ext == "csv":
        return _extract_csv_text(file_path), {"kind": "csv"}
    # .txt and other plain text files
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read(), {"kind": "text"}


def extract_text(file_path):
    """Extracts text from a PDF (normal or scanned), image, DOCX, CSV or TXT file."""
    return extract_text_ex(file_path)[0]


CJK_CHUNK_SIZE = 250     # Chinese/Japanese/Korean: 1 character ~ 1-2 tokens, is liye chhota chunk
CJK_CHUNK_OVERLAP = 40
_CJK_RE = re.compile(r"[\u3040-\u30ff\u3400-\u4dbf\u4e00-\u9fff\uac00-\ud7af\uf900-\ufaff]")


def _cjk_ratio(text):
    sample = text[:5000]
    if not sample:
        return 0.0
    return len(_CJK_RE.findall(sample)) / len(sample)


def chunk_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    """Splits text into overlapping chunks so context isn't lost at boundaries.

    Agar text zyada-tar Chinese/Japanese/Korean hai (aur caller ne default
    sizes hi use kiye), to chhote chunks (250 chars) banate hain -- warna 500
    CJK characters embedding model ki token limit se bahar chale jate hain
    aur chunk ka aakhri hissa search me shamil hi nahi hota."""
    text = text.strip()
    if (chunk_size, overlap) == (CHUNK_SIZE, CHUNK_OVERLAP) and _cjk_ratio(text) > 0.3:
        chunk_size, overlap = CJK_CHUNK_SIZE, CJK_CHUNK_OVERLAP
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += chunk_size - overlap
    return [c.strip() for c in chunks if c.strip()]


def _paths_for_user(user_id):
    user_dir = os.path.join(VECTORSTORE_DIR, str(user_id))
    os.makedirs(user_dir, exist_ok=True)
    index_path = os.path.join(user_dir, "index.faiss")
    meta_path = os.path.join(user_dir, "meta.pkl")
    return index_path, meta_path


def _model_marker_path(user_id):
    return os.path.join(VECTORSTORE_DIR, str(user_id), "embed_model.txt")


def _full_text_path(user_id):
    return os.path.join(VECTORSTORE_DIR, str(user_id), "full_text.txt")


def load_full_document_text(user_id):
    """
    Poora extracted text (chunk hone SE PEHLE, jaisa document se nikla tha)
    -- "poora document as-is dedo" jaisi requests ke liye. Har naye upload
    par overwrite hota hai (index bhi waisi hi replace hoti hai), is liye
    hamesha abhi ke active document se match karta hai.
    Return: (text, filename) ya (None, None) agar kuch upload hi nahi hua.
    """
    index_path, meta_path = _paths_for_user(user_id)
    text_path = _full_text_path(user_id)
    if not (os.path.exists(index_path) and os.path.exists(text_path)):
        return None, None
    try:
        with open(meta_path, "rb") as f:
            metadata = pickle.load(f)
        filename = metadata[0]["source"] if metadata else None
        with open(text_path, "r", encoding="utf-8") as f:
            return f.read(), filename
    except (OSError, EOFError, pickle.PickleError, IndexError, KeyError):
        return None, None


def _index_matches_current_model(user_id):
    """
    True sirf tab jab is user ka index usi embedding model se bana ho jo abhi
    load hai. Model badalne se pehle bane indexes ke vectors naye model ke
    query vectors se compare nahi ho sakte (dimension same ho tab bhi) -- unhe
    ignore karte hain taake bekaar/ghalat retrieval na ho.
    """
    try:
        with open(_model_marker_path(user_id), "r", encoding="utf-8") as f:
            return f.read().strip() == EMBED_MODEL_NAME
    except OSError:
        return False


def build_or_update_index(user_id, file_path, filename, session_id=None):
    """
    Processes a new document and builds a fresh FAISS index for the user,
    replacing any previously uploaded document. Each user has their own
    separate index (uploads stay private), and only the most recently
    uploaded document is searchable at any given time.

    Langfuse par ye poora pipeline ek "document-ingest" trace banata hai
    jiske andar chaar spans hain: extract-text, chunking, embedding,
    index-write. Is se dashboard par saaf nazar aata hai ke 20-page PDF
    upload karne me waqt kahan lag raha hai -- text nikalne me, embeddings
    banane me, ya index likhne me.
    """
    langfuse = _get_langfuse_client()
    file_type = (os.path.splitext(filename)[1].lstrip(".").lower() or "unknown")

    with _observe(
        langfuse, "document-ingest", as_type="span",
        input={"filename": filename, "fileType": file_type},
    ) as root_span:
        with _attributes(
            langfuse,
            user_id=f"user_{user_id}",
            session_id=session_id,
            metadata={"mode": "ingest"},
        ):
            _emit_event(langfuse, "ingest-request", metadata={
                "filename": filename,
                "fileType": file_type,
            })

            # ---- 1. Text extraction ----
            with _observe(langfuse, "extract-text", as_type="span",
                          input={"filename": filename, "fileType": file_type}) as span:
                text, extract_info = extract_text_ex(file_path)
                _update(span, output={
                    "characters": len(text),
                    "isEmpty": not text.strip(),
                    **{k: v for k, v in extract_info.items() if k != "kind"},
                })

            # ---- 2. Chunking ----
            with _observe(langfuse, "chunking", as_type="span",
                          input={"chunkSize": CHUNK_SIZE, "overlap": CHUNK_OVERLAP}) as span:
                chunks = chunk_text(text)
                _update(span, output={
                    "chunkCount": len(chunks),
                    "avgChunkChars": (
                        round(sum(len(c) for c in chunks) / len(chunks)) if chunks else 0
                    ),
                })

            if not chunks:
                # Khaali file, ya OCR ko bhi koi text nahi mila -- yahin ruk jate hain.
                _emit_event(langfuse, "ingest-skipped", metadata={
                    "reason": "no extractable text (empty file, or OCR found no text)",
                    "filename": filename,
                })
                _tag_trace(
                    langfuse,
                    tags=[TRACE_APP_TAG, "ingest", "empty-document"],
                    output={"chunkCount": 0, "indexed": False},
                )
                _update(root_span, output={"chunkCount": 0, "indexed": False})
                _flush(langfuse)
                return 0

            # ---- 3. Embedding ----
            with _observe(langfuse, "embedding", as_type="span",
                          input={"model": EMBED_MODEL_NAME, "chunkCount": len(chunks)}) as span:
                model = get_embed_model()
                embeddings = model.encode(
                    chunks, convert_to_numpy=True,
                    show_progress_bar=False, normalize_embeddings=True,
                )
                embeddings = embeddings.astype("float32")
                _update(span, output={
                    "vectors": int(embeddings.shape[0]),
                    "dimensions": int(embeddings.shape[1]),
                })

            # ---- 4. Index write ----
            with _observe(langfuse, "index-write", as_type="span") as span:
                index_path, meta_path = _paths_for_user(user_id)

                # Start a fresh index for every new upload, discarding any
                # previously uploaded document's chunks. This ensures questions
                # are answered only from the most recently uploaded document.
                index = faiss.IndexFlatL2(embeddings.shape[1])
                index.add(embeddings)
                metadata = [{"text": chunk, "source": filename} for chunk in chunks]

                faiss.write_index(index, index_path)
                with open(meta_path, "wb") as f:
                    pickle.dump(metadata, f)
                with open(_model_marker_path(user_id), "w", encoding="utf-8") as f:
                    f.write(EMBED_MODEL_NAME)
                # "poora document as-is dedo" jaisi requests ke liye asal text
                # (chunking se pehle) alag se save karte hain.
                with open(_full_text_path(user_id), "w", encoding="utf-8") as f:
                    f.write(text)

                _update(span, output={
                    "vectorsInIndex": int(index.ntotal),
                    "replacedPreviousIndex": True,
                })

            _emit_event(langfuse, "ingest-completed", metadata={
                "filename": filename,
                "chunkCount": len(chunks),
            })
            _tag_trace(
                langfuse,
                tags=[TRACE_APP_TAG, "ingest", file_type],
                output={"filename": filename, "chunkCount": len(chunks), "indexed": True},
            )

        _update(root_span, output={
            "filename": filename,
            "chunkCount": len(chunks),
            "indexed": True,
        })

    _flush(langfuse)
    return len(chunks)


def retrieve_relevant_chunks(user_id, query, top_k=4, query_vec=None):
    """
    Fetches the top-k chunks most relevant to the query from the user's index.

    `query_vec`: pehle se bana hua query embedding. answer_question() ise
    pass karti hai taake embedding apne alag Langfuse span me ho. Na diya
    jaye to yahin bana liya jata hai (purane callers waise hi chalte rehte
    hain).

    Note: we deliberately do NOT hard-filter these by embedding distance.
    Short/meta-phrased questions (e.g. "what is the candidate's name in the
    document") often score a weak embedding-similarity distance against the
    actual resume text even though they ARE about the document -- and that
    score range overlaps with genuinely unrelated questions. A fixed
    threshold can't reliably tell the two apart. Instead, all top-k chunks
    are always handed to the LLM, and the "rag-answer" prompt itself judges
    whether the content is actually relevant (see generate_answer() and the
    SOURCE_USED marker it parses from the model's answer).
    """
    index_path, meta_path = _paths_for_user(user_id)
    if not os.path.exists(index_path):
        return []
    if not _index_matches_current_model(user_id):
        print(f"[RAG] user {user_id}: index was built with a different embedding model -- ignoring it, please re-upload the document.")
        return []

    index = faiss.read_index(index_path)
    with open(meta_path, "rb") as f:
        metadata = pickle.load(f)

    if query_vec is None:
        query_vec = embed_query(query)

    k = min(top_k, index.ntotal)
    if k == 0:
        return []
    distances, indices = index.search(query_vec, k)

    # Har chunk ke sath uski distance aur rank bhi rakh lete hain -- ye
    # Langfuse ke retrieval span me chala jata hai, jahan se andaza hota
    # hai ke match kitna mazboot tha.
    results = []
    for rank, (idx, distance) in enumerate(zip(indices[0], distances[0]), start=1):
        if 0 <= idx < len(metadata):
            chunk = dict(metadata[idx])
            chunk["distance"] = float(distance)
            chunk["rank"] = rank
            results.append(chunk)

    if len(distances[0]) > 0:
        print(f"[RAG] query={query!r} best_distance={distances[0][0]:.3f} (informational only, not filtered)")

    return results